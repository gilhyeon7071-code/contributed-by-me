"""Trace unresolved current-day entry origins without changing trading state."""
from __future__ import annotations

import csv
import datetime as dt
import json
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
LOG_DIR = ROOT / "2_Logs"
CODES = {"003280", "058970", "195940", "087010", "328380"}


def _read_csv(path: Path) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return [{str(k): str(v) for k, v in row.items()} for row in csv.DictReader(f)]


def _read_orders_exec(path: Path) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "") for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: List[Dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[i]: ("" if value is None else str(value)) for i, value in enumerate(row)}
        if item.get("code") in CODES or item.get("종목코드") in CODES:
            rows.append(item)
    return rows


def _f(value: Any) -> float:
    try:
        return float(str(value).strip())
    except Exception:
        return 0.0


def _now() -> str:
    kst = dt.timezone(dt.timedelta(hours=9))
    return dt.datetime.now(tz=kst).isoformat(timespec="seconds")


def _write_csv(path: Path, rows: List[Dict[str, Any]]) -> None:
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    post = [
        row
        for row in _read_csv(LOG_DIR / "post_entry_learning_latest.csv")
        if row.get("code") in CODES and row.get("entry_origin") == "FILL_OBSERVED_UNRESOLVED"
    ]
    plan = [row for row in _read_csv(LOG_DIR / "candidate_action_plan_latest.csv") if row.get("code") in CODES]
    follow = [row for row in _read_csv(LOG_DIR / "candidate_action_followup_latest.csv") if row.get("code") in CODES]
    recheck = [
        row for row in _read_csv(LOG_DIR / "candidate_action_recheck_queue_latest.csv") if row.get("code") in CODES
    ]
    fills = [
        row
        for row in _read_csv(ROOT / "paper" / "fills.csv")
        if row.get("code") in CODES
        and row.get("side") == "BUY"
        and (row.get("datetime", "").startswith("20260612") or row.get("ymd") == "20260612")
    ]
    broker = [
        row
        for row in _read_csv(ROOT / "paper" / "orders_20260612_broker_submit_mock.csv")
        if row.get("code") in CODES and row.get("side") == "BUY"
    ]
    orders_exec_path = ROOT / "paper" / "orders_20260612_exec.xlsx"
    orders_exec = _read_orders_exec(orders_exec_path)

    rows: List[Dict[str, Any]] = []
    for code in sorted(CODES):
        post_row = next((row for row in post if row.get("code") == code), {})
        plan_rows = [row for row in plan if row.get("code") == code]
        exec_plan = next((row for row in plan_rows if row.get("next_action") == "EXECUTED_BUY"), {})
        non_exec_rows = [row for row in plan_rows if row.get("fill_provenance") == "SAME_CODE_NON_EXECUTION_ROW"]
        fill_rows = [row for row in fills if row.get("code") == code]
        broker_rows = [row for row in broker if row.get("code") == code]
        exec_rows = [row for row in orders_exec if row.get("code") == code or row.get("종목코드") == code]
        broker_status_counts = dict(Counter(row.get("dispatch_status", "") for row in broker_rows))

        if (
            exec_plan.get("fill_provenance") == "FILL_OBSERVED_NO_CURRENT_TRADABLE_ROW"
            and exec_plan.get("reason") == "not_execution_pool"
        ):
            root_cause_class = "FILL_OBSERVED_AFTER_NOT_EXECUTION_POOL_NO_CURRENT_TRADABLE_ROW"
        else:
            root_cause_class = "UNCLASSIFIED_NEEDS_DEEP_TRACE"

        ret = _f(post_row.get("unrealized_return_pct"))
        if ret < 0 or post_row.get("exit_reason"):
            risk_bucket = "LOSS_OR_STOP"
        elif ret > 0:
            risk_bucket = "POSITIVE_OPEN"
        else:
            risk_bucket = "FLAT_OPEN"

        rows.append(
            {
                "code": code,
                "name": post_row.get("name", ""),
                "entry_ts": post_row.get("entry_ts", ""),
                "order_id": post_row.get("order_id", ""),
                "entry_price": post_row.get("entry_price", ""),
                "open_qty": post_row.get("open_qty", ""),
                "unrealized_return_pct": post_row.get("unrealized_return_pct", ""),
                "exit_reason": post_row.get("exit_reason", ""),
                "learning_bucket": post_row.get("learning_bucket", ""),
                "learning_policy_action": post_row.get("learning_policy_action", ""),
                "exec_plan_next_action": exec_plan.get("next_action", ""),
                "exec_plan_action_state": exec_plan.get("action_state", ""),
                "exec_plan_reason": exec_plan.get("reason", ""),
                "exec_plan_trading_allowed": exec_plan.get("trading_allowed", ""),
                "exec_plan_fill_provenance": exec_plan.get("fill_provenance", ""),
                "same_code_non_execution_rows": len(non_exec_rows),
                "same_code_non_execution_reasons": "|".join(
                    sorted({row.get("action_reason", "") for row in non_exec_rows if row.get("action_reason")})
                ),
                "followup_reasons": "|".join(
                    sorted({row.get("reason", "") for row in follow if row.get("code") == code and row.get("reason")})
                ),
                "recheck_reasons": "|".join(
                    sorted({row.get("reason", "") for row in recheck if row.get("code") == code and row.get("reason")})
                ),
                "fills_buy_rows_20260612": len(fill_rows),
                "broker_buy_rows_20260612": len(broker_rows),
                "broker_dispatch_status_counts": json.dumps(broker_status_counts, ensure_ascii=False, sort_keys=True),
                "orders_exec_rows_20260612": len(exec_rows),
                "root_cause_class": root_cause_class,
                "risk_bucket": risk_bucket,
            }
        )

    summary = {
        "traced_codes": len(rows),
        "root_cause_class_counts": dict(Counter(row["root_cause_class"] for row in rows)),
        "risk_bucket_counts": dict(Counter(row["risk_bucket"] for row in rows)),
        "same_code_non_execution_codes": [
            row["code"] for row in rows if int(row["same_code_non_execution_rows"]) > 0
        ],
        "no_same_code_non_execution_codes": [
            row["code"] for row in rows if int(row["same_code_non_execution_rows"]) == 0
        ],
        "broker_only_skip_already_filled_codes": [
            row["code"]
            for row in rows
            if row["broker_buy_rows_20260612"] > 0
            and set(json.loads(row["broker_dispatch_status_counts"]).keys()) == {"SKIP_ALREADY_PAPER_FILLED"}
        ],
        "missing_from_broker_submit_codes": [row["code"] for row in rows if row["broker_buy_rows_20260612"] == 0],
    }

    out_json = LOG_DIR / "unresolved_entry_origin_trace_latest.json"
    out_csv = LOG_DIR / "unresolved_entry_origin_trace_latest.csv"
    payload = {
        "generated_at": _now(),
        "schema_version": "unresolved_entry_origin_trace_v2",
        "trading_effect": False,
        "policy_effect": False,
        "policy_change_applied": False,
        "scope": "FILL_OBSERVED_UNRESOLVED current-day buys only; CSV/XLSX trace",
        "source_artifacts": {
            "post_entry_learning_csv": str(LOG_DIR / "post_entry_learning_latest.csv"),
            "candidate_action_plan_csv": str(LOG_DIR / "candidate_action_plan_latest.csv"),
            "candidate_action_followup_csv": str(LOG_DIR / "candidate_action_followup_latest.csv"),
            "candidate_action_recheck_queue_csv": str(LOG_DIR / "candidate_action_recheck_queue_latest.csv"),
            "fills_csv": str(ROOT / "paper" / "fills.csv"),
            "broker_submit_mock_csv": str(ROOT / "paper" / "orders_20260612_broker_submit_mock.csv"),
            "orders_exec_xlsx": str(orders_exec_path),
        },
        "summary": summary,
        "rows": rows,
        "interpretation": {
            "finding": (
                "All five unresolved entries are observed fills whose current action-plan "
                "EXECUTED_BUY rows are blocked/not_execution_pool and marked "
                "FILL_OBSERVED_NO_CURRENT_TRADABLE_ROW."
            ),
            "split": (
                "003280/058970/328380 also have same-code non-execution recheck rows; "
                "087010/195940 do not in current plan."
            ),
            "boundary": "Read-only trace only; no threshold, gate, order, fill, ledger, or policy change.",
            "next_root_cause_target": (
                "Find the earlier intraday candidate snapshot/order creation path that created "
                "PAPER_BUY before current not_execution_pool state."
            ),
        },
    }
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_csv(out_csv, rows)
    print(json.dumps({"status": "OK", "rows": len(rows), "out_json": str(out_json)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
